#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MONTA candidatas/ITALY-SOURCE-GAP-BY-TOOL-2026-09-14.xlsx

A FOLHA CENTRAL E' SOURCE_TO_TOOL, E E' A UNICA QUE REPETE FONTE
----------------------------------------------------------------
Uma fonte serve varias ferramentas. NEW_SOURCES tem uma linha por fonte;
SOURCE_TO_TOOL tem uma linha por LIGACAO. Trocar as duas inflacionaria a
entrega: 37 fontes tornam-se 109 linhas de ligacao, e quem contasse linhas
diria que a missao achou 109 fontes.

Por isso cada folha declara no proprio cabecalho o que uma linha e'.
"""

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "candidatas"))
TRAB = Path("C:/Users/London1/AppData/Local/Temp/sintonia-gap")

from italy_gap_needs import NECESSIDADES, FERRAMENTAS          # noqa: E402
from italy_gap_auditoria import AUDITORIA                      # noqa: E402
from italy_gap_novas_fontes import TODAS as NOVAS, ROTAS_MORTAS  # noqa: E402
from italy_gap_buscas import BUSCAS                            # noqa: E402

from openpyxl import Workbook                                  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill       # noqa: E402
from openpyxl.utils import get_column_letter                   # noqa: E402

SAIDA = RAIZ / "candidatas" / "ITALY-SOURCE-GAP-BY-TOOL-2026-09-14.xlsx"
COBERTURA = json.loads((TRAB / "COBERTURA.json").read_text(encoding="utf-8"))
MATRIZES = json.loads((TRAB / "MATRIZES.json").read_text(encoding="utf-8"))

CAB = PatternFill("solid", fgColor="1F3864")
NOTA = PatternFill("solid", fgColor="FFF2CC")
COR = {"CRITICAL": "F4CCCC", "HIGH": "FCE5CD", "MEDIUM": "FFF2CC",
       "LOW": "D9EAD3"}


def folha(wb, nome, explica, colunas, linhas):
    ws = wb.create_sheet(nome)
    ws["A1"] = explica
    ws["A1"].fill = NOTA
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1,
                   end_column=max(len(colunas), 2))
    ws.row_dimensions[1].height = 44
    for j, c in enumerate(colunas, 1):
        cel = ws.cell(row=2, column=j, value=c)
        cel.fill = CAB
        cel.font = Font(color="FFFFFF", bold=True)
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    for i, lin in enumerate(linhas, 3):
        for j, c in enumerate(colunas, 1):
            v = lin.get(c, "")
            cel = ws.cell(row=i, column=j,
                          value=v if isinstance(v, (int, float)) else str(v))
            cel.alignment = Alignment(wrap_text=True, vertical="top")
        sev = str(lin.get("GAP_SEVERITY", ""))
        if sev in COR:
            ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=COR[sev])
    for j, c in enumerate(colunas, 1):
        larg = max([len(c)] + [min(len(str(l.get(c, ""))), 78) for l in linhas]) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(larg, 10), 62)
    ws.freeze_panes = "A3"
    return ws


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    wb = Workbook()
    wb.remove(wb.active)

    # ── 1 · TOOL_RAW_NEEDS ────────────────────────────────────────────────
    l1 = []
    for n in NECESSIDADES:
        f = FERRAMENTAS[n["TOOL"]]
        l1.append({
            "TOOL": n["TOOL"], "FERRAMENTA_NO_PORTAL": f["nome_it"],
            "PECA_NO_MAPA": f["id_mapa"],
            "PERGUNTA_QUE_A_FERRAMENTA_RESPONDE": f["pergunta"],
            "RAW_NEED": n["RAW_NEED"], "REQUIRED_FIELDS": n["REQUIRED_FIELDS"],
            "FRESHNESS": n["FRESHNESS"], "GEO": n["GEO"],
            "TEMPORAL": n["TEMPORAL"], "TERRITORIO": " ".join(n["TERR"]),
            "ANCORA": n["ANCORA"], "PROVA_DA_NECESSIDADE": n["PROVA"],
            "O_QUE_O_CASCO_TEM_HOJE": n["MEDIDO"]})
    folha(wb, "TOOL_RAW_NEEDS",
          "UMA LINHA = UMA MATERIA-PRIMA que uma ferramenta precisa ANTES da "
          "Intelligence. As 9 ferramentas foram MEDIDAS no portal (zona "
          "Z-TELAS), nao copiadas do pedido: Archivio e Registro delle fonti "
          "ficaram fora porque nao sao Intelligence. ANCORA=CONTRATO quer "
          "dizer que a necessidade vem de um ficheiro de contrato do bloco; "
          "ANCORA=DADO_MEDIDO quer dizer que vem da leitura do dado real.",
          list(l1[0]), l1)

    # ── 2 · CURRENT_COVERAGE ──────────────────────────────────────────────
    l2 = []
    for r in COBERTURA:
        k = f'{r["TOOL"]} · {r["RAW_NEED"]}'
        ver, quem, porque = AUDITORIA[k]
        l2.append({
            "TOOL": r["TOOL"], "RAW_NEED": r["RAW_NEED"],
            "GAP_SEVERITY": r["GAP_SEVERITY"],
            "CASCO_TEM_O_DADO": r["CASCO_TEM_O_DADO"],
            "ENCHIMENTO_MEDIDO": r["ENCHIMENTO_NO_CASCO"],
            "MEDICAO": r["MEDIDO"],
            "VEREDITO_AUTOMATICO": r["COBERTURA"],
            "DONOS_DISTINTOS_PROPOSTOS": r["DONOS_FORTES"],
            "VEREDITO_A_MAO": ver,
            "QUEM_SUSTENTA": quem,
            "POR_QUE": porque,
            "MATCHER": r["MATCHER"], "MATCHER_PORQUE": r["MATCHER_PORQUE"]})
    folha(wb, "CURRENT_COVERAGE",
          "UMA LINHA = UMA NECESSIDADE, com TRES contas que nao se misturam: "
          "(1) o casco TEM o dado hoje? — medido campo a campo; (2) o que o "
          "casamento automatico por palavra propos; (3) o que a leitura a mao "
          "decidiu. O automatico NUNCA escreve STRONG: ele so prova que a "
          "fonte FALA do assunto. Quatro propostas suas foram derrubadas a "
          "mao — entre elas um registo frances para fenologia e uma clinica "
          "veterinaria para voz de campo.",
          list(l2[0]), l2)

    # ── 3 · CRITICAL_GAPS ─────────────────────────────────────────────────
    l3 = []
    for r in COBERTURA:
        if r["GAP_SEVERITY"] not in ("CRITICAL", "HIGH"):
            continue
        k = f'{r["TOOL"]} · {r["RAW_NEED"]}'
        ver, quem, porque = AUDITORIA[k]
        fecha = {"STRONG": "SIM — ha fonte nomeada",
                 "STRONG_MAS_TROCA_A_VOZ": "PARCIAL — enche a coluna, troca o sujeito",
                 "WEAK": "NAO AINDA — fonte plausivel, campo nao visto",
                 "NO": "NAO — e a razao nao e' tecnica",
                 "NAO_E_GAP_DE_FONTE": "NAO SE FECHA COM FONTE",
                 "HERDADO": "= ferramenta de origem"}[ver]
        l3.append({
            "GAP_SEVERITY": r["GAP_SEVERITY"], "TOOL": r["TOOL"],
            "RAW_NEED": r["RAW_NEED"],
            "O_QUE_A_FERRAMENTA_PERDE_SEM_ISTO": FERRAMENTAS[r["TOOL"]]["pergunta"],
            "MEDICAO_QUE_PROVA_O_GAP": r["MEDIDO"],
            "FONTE_NOVA_FECHA_ISTO": fecha,
            "VEREDITO_A_MAO": ver, "QUEM_SUSTENTA": quem, "POR_QUE": porque})
    l3.sort(key=lambda x: (x["GAP_SEVERITY"] != "CRITICAL", x["TOOL"]))
    folha(wb, "CRITICAL_GAPS",
          "UMA LINHA = UM GAP CRITICAL ou HIGH. A gravidade sai do CAMPO "
          "MEDIDO no casco, nunca da contagem de candidatas — ter candidata e' "
          "promessa, ter campo e' capacidade. Leia a coluna "
          "FONTE_NOVA_FECHA_ISTO: nem todo gap fecha com fonte. Tres classes "
          "nao fecham — join interno, extracao de PDF, e dado proprietario da "
          "casa.",
          list(l3[0]), l3)

    # ── 4 · NEW_SOURCES ───────────────────────────────────────────────────
    l4 = []
    for f in NOVAS:
        l4.append({
            "IDENTIDADE": f["NOME"], "OWNER": f["DONO"], "URL": f["URL"],
            "CLASSE": f["CLASSE"],
            "O_QUE_PRODUZ": f["PRODUZ"], "EXEMPLO_REAL": f["EXEMPLO_REAL"],
            "QUANTAS_NECESSIDADES_ALIMENTA": len(f["ALIMENTA"]),
            "QUE_FERRAMENTAS_SERVE": " · ".join(sorted({a.split(" · ")[0]
                                                        for a in f["ALIMENTA"]})),
            "AMBITO_GEOGRAFICO": f["AMBITO"],
            "TERRITORIO": " ".join(f["TERR"]),
            "RECURRING_POTENTIAL": f["RECORRENCIA"],
            "EVIDENCIA": f["EVIDENCIA"],
            "LIMITE_DECLARADO": f["LIMITE"]})
    l4.sort(key=lambda x: (x["CLASSE"], -x["QUANTAS_NECESSIDADES_ALIMENTA"]))
    folha(wb, "NEW_SOURCES",
          "UMA LINHA = UMA FONTE, contada UMA VEZ mesmo quando serve varias "
          "ferramentas. Toda linha nasceu de um buraco medido, nunca de um "
          "tema. Classe A exige as dez coisas (identidade, dono, rota, "
          "exemplo real, o que produz, que necessidade alimenta, que "
          "ferramenta serve, ambito, recorrencia, evidencia); falta uma e e' "
          "B. ⚠️ NENHUMA destas rotas foi aberta por sonda nesta missao: a "
          "evidencia e' de descoberta, e rota descrita nao e' rota lida.",
          list(l4[0]), l4)

    # ── 5 · SOURCE_TO_TOOL  (a folha central) ─────────────────────────────
    l5 = []
    for f in NOVAS:
        for a in f["ALIMENTA"]:
            tool, need = a.split(" · ", 1)
            alvo = next((r for r in COBERTURA
                         if r["TOOL"] == tool and r["RAW_NEED"] == need), None)
            ver = AUDITORIA[a][0]
            l5.append({
                "SOURCE": f["NOME"], "OWNER": f["DONO"], "URL": f["URL"],
                "CLASSE_DA_FONTE": f["CLASSE"],
                "TOOL": tool, "FERRAMENTA_NO_PORTAL": FERRAMENTAS[tool]["nome_it"],
                "RAW_NEED_QUE_ALIMENTA": need,
                "GAP_SEVERITY_DA_NECESSIDADE": alvo["GAP_SEVERITY"] if alvo else "?",
                "VEREDITO_A_MAO_DA_NECESSIDADE": ver,
                "AMBITO": f["AMBITO"], "RECORRENCIA": f["RECORRENCIA"]})
    l5.sort(key=lambda x: (x["TOOL"], x["SOURCE"]))
    folha(wb, "SOURCE_TO_TOOL",
          "A FOLHA CENTRAL. UMA LINHA = UMA LIGACAO fonte→necessidade, nao "
          "uma fonte. 37 fontes produzem 109 ligacoes: quem contar linhas "
          "aqui e disser «109 fontes novas» conta a mesma fonte ate oito "
          "vezes. «Uma SOURCE, varios usos» le-se nesta folha e em nenhuma "
          "outra.",
          list(l5[0]), l5)

    # ── 6 · RADAR_INPUTS ──────────────────────────────────────────────────
    folha(wb, "RADAR_INPUTS",
          "O OPPORTUNITY RADAR NAO TEM FAMILIA PROPRIA DE FONTES. E' um "
          "cruzamento das outras camadas, e por isso aqui nao ha lista de "
          "fontes — ha matriz de ENTRADAS, cada uma a apontar para a camada "
          "que a produz. Procurar «fontes para o Radar» contaria a mesma "
          "fonte duas vezes. Regra dura: o Radar nunca e' melhor do que a sua "
          "entrada mais fraca.",
          list(MATRIZES["RADAR"][0]), MATRIZES["RADAR"])

    # ── 7 · REGION_GAPS ───────────────────────────────────────────────────
    folha(wb, "REGION_GAPS",
          "UMA LINHA = UMA DAS 20 REGIOES. A pergunta nao e' «quantas fontes "
          "tem esta regiao» — e' «ha fonte que enche o buraco CRITICO da "
          "fenologia aqui». ⚠️ FONTE NACIONAL NAO CONTA COMO COBERTURA "
          "REGIONAL, e esta em coluna separada: a primeira versao desta folha "
          "contava-a e dizia 20 de 20 cobertas, incluindo Molise e Valle "
          f"d'Aosta. As {len(MATRIZES['REGION_NACIONAIS'])} nacionais: "
          f"{' · '.join(MATRIZES['REGION_NACIONAIS'])}.",
          list(MATRIZES["REGION"][0]), MATRIZES["REGION"])

    # ── 8 · CROP_GAPS ─────────────────────────────────────────────────────
    folha(wb, "CROP_GAPS",
          "UMA LINHA = UMA DAS 10 CULTURAS das 29 janelas canonicas. A "
          "contagem procura o nome ITALIANO da cultura nos campos em italiano "
          "da fonte (nome, exemplo, dono) — nunca na descricao em portugues. "
          "⚠️ «mais» (milho em italiano) foi retirado das palavras-chave: "
          "colide com «mais» em portugues e produzia 6 falsos. As cinco "
          "culturas com zero nao provam ausencia de fonte em Italia — provam "
          "que a busca foi a vite, ao olivo, aos agrumes e aos frutos de pomo.",
          list(MATRIZES["CROP"][0]), MATRIZES["CROP"])

    # ── 9 · SOCIAL_GAPS ───────────────────────────────────────────────────
    folha(wb, "SOCIAL_GAPS",
          "UMA LINHA = UM CAMPO DA VOZ DE CAMPO, com o enchimento real. Nao "
          "e' lista de canais: e' a conta que explica por que o Radar nao "
          "consegue cruzar voz com janela. Tres campos a zero de 17, e sao "
          "exatamente os tres que o cruzamento precisa.",
          list(MATRIZES["SOCIAL"][0]), MATRIZES["SOCIAL"])

    # ── 10 · PEOPLE_GAPS ──────────────────────────────────────────────────
    folha(wb, "PEOPLE_GAPS",
          "UMA LINHA = UM PAPEL, com as pessoas nomeadas que a missao achou. "
          "Agrupado por PAPEL e nunca por alcance: seguidores nao sao "
          "autoridade, e isto nao e' ranking de especialistas. A linha mais "
          "importante e' a ultima, com QUANTAS=0.",
          list(MATRIZES["PEOPLE"][0]), MATRIZES["PEOPLE"])

    # ── 11 · DEAD_ROUTES ──────────────────────────────────────────────────
    l11 = [{"ROTA_NO_ACERVO": d["ROTA"], "ESTADO": d["ESTADO"],
            "SUBSTITUTA": d["SUBSTITUTA"],
            "POR_QUE_IMPORTA": d["PORQUE_IMPORTA"]} for d in ROTAS_MORTAS]
    folha(wb, "DEAD_ROUTES",
          "ACHADO QUE A MISSAO NAO PROCURAVA. UMA LINHA = UMA ROTA DO ACERVO "
          "QUE MORREU OU MUDOU DE REGRA. Rota encerrada raramente devolve "
          "erro limpo — devolve pagina. Coleta silenciosa de nada e' pior do "
          "que coleta falhada, porque ninguem investiga o que parece ter "
          "corrido bem.",
          list(l11[0]), l11)

    # ── 12 · SEARCH_LOG ───────────────────────────────────────────────────
    l12 = [{"N": i, "GAP_QUE_GUIOU_A_BUSCA": b["GAP"],
            "CONSULTA": b["CONSULTA"],
            "O_QUE_DEVOLVEU": b["DEVOLVEU"],
            "FONTES_NOVAS_DAQUI": b["FONTES"],
            "ACHADO_QUE_MUDOU_A_RESPOSTA": b["VIROU"]}
           for i, b in enumerate(BUSCAS, 1)]
    folha(wb, "SEARCH_LOG",
          "UMA LINHA = UMA BUSCA. Toda consulta parte de um GAP MEDIDO, nunca "
          "de um tema: nao «fontes agricolas da Sicilia», mas «Sicilia + vite "
          "+ bollettino fitosanitario + fenologia». A ultima coluna guarda as "
          "buscas que mudaram a RESPOSTA em vez de acrescentar fonte — foram "
          "essas as mais uteis.",
          list(l12[0]), l12)

    # ── 13 · README ───────────────────────────────────────────────────────
    sev = Counter(r["GAP_SEVERITY"] for r in COBERTURA)
    ver = Counter(v[0] for v in AUDITORIA.values())
    est = Counter(l["ESTADO"] for l in MATRIZES["REGION"])
    l13 = [
     {"O_QUE": "A PERGUNTA DA MISSAO", "RESPOSTA":
      "Que informacao BRUTA precisa existir ANTES da Intelligence para cada "
      "ferramenta funcionar, e quem em Italia produz essa informacao."},
     {"O_QUE": "FERRAMENTAS REAIS MEDIDAS", "RESPOSTA":
      f"{len(FERRAMENTAS)} ferramentas de Intelligence, medidas na zona "
      "Z-TELAS do System Map. 11 telas existem; Archivio e Registro delle "
      "fonti nao sao Intelligence e ficaram fora. A lista do pedido nao foi "
      "usada como verdade."},
     {"O_QUE": "MATERIAS-PRIMAS MAPEADAS", "RESPOSTA":
      f"{len(NECESSIDADES)} necessidades. Gravidade: "
      + " · ".join(f"{k}={sev[k]}" for k in ("CRITICAL", "HIGH", "MEDIUM", "LOW"))},
     {"O_QUE": "VEREDITO A MAO (72 de 72)", "RESPOSTA":
      " · ".join(f"{k}={v}" for k, v in sorted(ver.items()))},
     {"O_QUE": "FONTES NOVAS", "RESPOSTA":
      f"{len(NOVAS)} fontes, {sum(len(f['ALIMENTA']) for f in NOVAS)} ligacoes "
      f"a {len({a for f in NOVAS for a in f['ALIMENTA']})} necessidades "
      "distintas. Classe: "
      + " · ".join(f"{k}={v}" for k, v in
                   sorted(Counter(f["CLASSE"] for f in NOVAS).items()))},
     {"O_QUE": "REGIOES", "RESPOSTA":
      " · ".join(f"{k}={v}" for k, v in sorted(est.items()))
      + ". Fonte nacional contada a parte."},
     {"O_QUE": "O QUE NAO SE FEZ", "RESPOSTA":
      "COLLECTION_RUNS_CREATED = 0 · RAW_CREATED = 0 · SOURCE_ID_CREATED = 0 "
      "· DOCUMENT_ID_CREATED = 0 · WAITING_ROOM_DELTA = 0 · "
      "INTELLIGENCE_RUNS = 0 · PORTAL_CHANGED = NO · ATLAS_CHANGED = NO. "
      "Nenhuma fonte foi registada; nenhuma taxonomia foi criada; nenhuma "
      "arquitetura nova foi desenhada."},
     {"O_QUE": "O QUE ISTO NAO PROVA", "RESPOSTA":
      "Nenhuma das 37 rotas novas foi aberta por sonda: a evidencia e' de "
      "busca dirigida. Nenhum exemplo de conteudo foi guardado. Nenhuma "
      "afiliacao de pessoa foi verificada em perfil institucional. Rota "
      "descrita nao e' rota lida, e modelo nao e' prova — a fonte real "
      "continua a ser a prova."},
     {"O_QUE": "GAPS QUE NAO SAO GAPS DE FONTE", "RESPOSTA":
      f"{ver['NAO_E_GAP_DE_FONTE']} necessidades nao se fecham com fonte "
      "nenhuma: join interno (produto×janela), procedencia nao escrita "
      "(SOURCE_IDS vazio), vocabulario e identidade internos, e o relato da "
      "rede comercial — que e' dado PROPRIO da casa. Coletar mais nao "
      "conserta nada disto."},
     {"O_QUE": "O GAP QUE E' LEGAL, NAO TECNICO", "RESPOSTA":
      "Os dados das provas GEP de registo sao reservados, propriedade de quem "
      "as encomenda, e nao sao publicados. ENSAIO DE CAMPO com resultado nao "
      "e' gap de fonte: e' gap de ACESSO, e resolve-se por contrato."},
     {"O_QUE": "COMO LER AS FOLHAS", "RESPOSTA":
      "NEW_SOURCES = uma linha por fonte. SOURCE_TO_TOOL = uma linha por "
      "LIGACAO (a mesma fonte repete-se ali, e so ali). CURRENT_COVERAGE tem "
      "o veredito automatico E o veredito a mao, lado a lado, de proposito: "
      "para se ver onde a maquina errou."},
     {"O_QUE": "DONO DA TAXONOMIA", "RESPOSTA":
      "docs/fontes/ATLAS-DE-FONTES-EAME.md, via _territorios.py. Todo codigo "
      "T desta pasta foi validado contra ele. T13 nao e' canonico e nao "
      "aparece aqui."},
    ]
    folha(wb, "README", "COMECE AQUI.", list(l13[0]), l13)

    wb.save(SAIDA)
    print(f"gravado: {SAIDA}")
    print(f"folhas: {len(wb.sheetnames)} — {', '.join(wb.sheetnames)}")
    for nome, n in (("TOOL_RAW_NEEDS", len(l1)), ("CURRENT_COVERAGE", len(l2)),
                    ("CRITICAL_GAPS", len(l3)), ("NEW_SOURCES", len(l4)),
                    ("SOURCE_TO_TOOL", len(l5)),
                    ("RADAR_INPUTS", len(MATRIZES["RADAR"])),
                    ("REGION_GAPS", len(MATRIZES["REGION"])),
                    ("CROP_GAPS", len(MATRIZES["CROP"])),
                    ("SOCIAL_GAPS", len(MATRIZES["SOCIAL"])),
                    ("PEOPLE_GAPS", len(MATRIZES["PEOPLE"])),
                    ("DEAD_ROUTES", len(l11)), ("SEARCH_LOG", len(l12)),
                    ("README", len(l13))):
        print(f"  {nome:18s} {n:4d} linhas")


if __name__ == "__main__":
    main()
